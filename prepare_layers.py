"""Build the editable web data used by the Friant research microsite.

Run with the same Python environment used for the research notebooks:

    python prepare_layers.py --package "/path/to/Friant_Drought_Research_Package"

The script never changes the research package. It writes simplified web layers,
raster overlays and JavaScript data files beside this website.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

import geopandas as gpd
import matplotlib
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from matplotlib.colors import BoundaryNorm, Normalize
from openpyxl import load_workbook
from rasterio.enums import Resampling
from rasterio.warp import calculate_default_transform, reproject, transform_bounds

matplotlib.use("Agg")

DEFAULT_PACKAGE_ROOT = Path(
    "/Users/samsundarsingh/Library/CloudStorage/OneDrive-IndianInstituteofScience/"
    "PostDocs/Friant_Drought_Research_Package"
)


def clean_number(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def canonical(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def write_js(path: Path, variable: str, payload) -> None:
    path.write_text(
        f"window.{variable} = "
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=clean_number)
        + ";\n",
        encoding="utf-8",
    )


def geojson_dict(frame: gpd.GeoDataFrame, simplify_m: float | None = None) -> dict:
    work = frame.copy()
    if work.crs is None:
        raise ValueError("GeoDataFrame has no coordinate reference system")
    if simplify_m:
        projected = work.to_crs(3310)
        projected.geometry = projected.geometry.simplify(simplify_m, preserve_topology=True)
        work = projected.to_crs(4326)
    else:
        work = work.to_crs(4326)
    work = work[work.geometry.notna() & ~work.geometry.is_empty].copy()
    for col in work.columns:
        if col == "geometry":
            continue
        if pd.api.types.is_datetime64_any_dtype(work[col]):
            work[col] = work[col].astype(str)
        elif pd.api.types.is_numeric_dtype(work[col]):
            work[col] = work[col].map(clean_number)
        else:
            work[col] = work[col].where(work[col].notna(), None)
    return json.loads(work.to_json(drop_id=True))


def workbook_table(path: Path, sheet_name: str) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(
        i
        for i, row in enumerate(rows)
        if row and ((row[0] is not None and "contractor" in str(row[0]).lower()) or row[0] == "year")
    )
    header = list(rows[header_idx])
    width = len(header)
    data = [list(row[:width]) for row in rows[header_idx + 1 :] if row and any(v is not None for v in row[:width])]
    return pd.DataFrame(data, columns=header)


def export_raster_overlay(src_path: Path, out_path: Path, cmap_name: str, alpha: int = 205,
                          discrete_bounds: list[float] | None = None) -> dict:
    with rasterio.open(src_path) as src:
        src_data = src.read(1, masked=True)
        dst_crs = "EPSG:4326"
        target_width = min(1500, src.width)
        scale = src.width / target_width
        target_height = max(1, int(src.height / scale))
        transform, width, height = calculate_default_transform(
            src.crs, dst_crs, src.width, src.height, *src.bounds,
            dst_width=target_width, dst_height=target_height,
        )
        dst = np.full((height, width), np.nan, dtype="float32")
        reproject(
            source=src_data.filled(np.nan),
            destination=dst,
            src_transform=src.transform,
            src_crs=src.crs,
            dst_transform=transform,
            dst_crs=dst_crs,
            resampling=Resampling.nearest if discrete_bounds else Resampling.bilinear,
            src_nodata=src.nodata,
            dst_nodata=np.nan,
        )
        valid = np.isfinite(dst)
        if not valid.any():
            raise ValueError(f"No valid cells in {src_path.name}")
        cmap = plt.get_cmap(cmap_name)
        if discrete_bounds:
            norm = BoundaryNorm(discrete_bounds, cmap.N)
        else:
            lo, hi = np.nanpercentile(dst[valid], [2, 98])
            if lo == hi:
                hi = lo + 1
            norm = Normalize(vmin=lo, vmax=hi, clip=True)
        rgba = (cmap(norm(dst)) * 255).astype("uint8")
        rgba[..., 3] = np.where(valid, alpha, 0).astype("uint8")
        plt.imsave(out_path, rgba)
        west, south, east, north = transform_bounds(src.crs, dst_crs, *src.bounds, densify_pts=21)
        return {
            "url": f"assets/{out_path.name}",
            "bounds": [[south, west], [north, east]],
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", type=Path, default=DEFAULT_PACKAGE_ROOT)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()

    package = args.package.expanduser().resolve()
    output = args.output.expanduser().resolve()
    assets = output / "assets"
    data_dir = output / "data"
    assets.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)

    gpkg = package / "02_Data/GIS/Friant_Division_All_Layers.gpkg"
    study = gpd.read_file(gpkg, layer="study_area")[["study_name", "geometry"]]
    canals = gpd.read_file(gpkg, layer="federal_canals")[["canal_mc", "geometry"]]
    turnouts = gpd.read_file(gpkg, layer="approximate_turnouts")[
        ["turnout", "milepost", "accuracy", "geometry"]
    ]
    rise = gpd.read_file(gpkg, layer="usbr_rise_locations")[["name", "rise_id", "geometry"]]

    wells = gpd.read_file(gpkg, layer="irrigation_well_records")[
        ["WCRNumber", "CountyName", "RecordType", "TotalCompletedDepth", "DateWorkEnded", "geometry"]
    ]
    wells["completed_year"] = pd.to_datetime(wells["DateWorkEnded"], unit="ms", errors="coerce").dt.year
    wells = wells.drop(columns="DateWorkEnded")

    monitors = gpd.read_file(gpkg, layer="groundwater_monitors")[
        ["id", "monitoring_location_name", "county_name", "site_type", "altitude", "geometry"]
    ]

    lulc = gpd.read_file(gpkg, layer="lulc_fields_2023")[
        ["map_group", "ACRES", "MAIN_CROP", "COUNTY", "geometry"]
    ]
    crop_summary = (
        lulc.groupby("map_group", dropna=False)
        .agg(field_count=("map_group", "size"), acres=("ACRES", "sum"))
        .reset_index()
        .sort_values("acres", ascending=False)
    )
    lulc_web = lulc.to_crs(3310)
    lulc_web.geometry = lulc_web.geometry.simplify(60, preserve_topology=True)
    lulc_web["field_count"] = 1
    lulc_web = lulc_web.dissolve(by="map_group", aggfunc={"ACRES": "sum", "field_count": "sum"}).reset_index()
    lulc_web = lulc_web.rename(columns={"ACRES": "acres"})

    groundwater_dir = package / "02_Data/Groundwater/Raw"
    geology = gpd.read_file(groundwater_dir / "CGS_Generalized_Rock_Types_Friant.gpkg")[
        ["GENERAL_LITHOLOGY", "AGE", "DESCRIPTION", "geometry"]
    ]
    geology = geology.dissolve(by=["GENERAL_LITHOLOGY", "AGE"], as_index=False)

    soils = gpd.read_file(groundwater_dir / "SSURGO_MapunitPoly_Friant.gpkg")[
        ["hydgrp", "geometry"]
    ]
    soils["hydgrp"] = soils["hydgrp"].fillna("Unclassified")
    soils = soils.dissolve(by="hydgrp", as_index=False)

    basins = gpd.read_file(groundwater_dir / "DWR_Bulletin118_Groundwater_Basins_Friant.gpkg")[
        ["Basin_Subbasin_Number", "Basin_Subbasin_Name", "Area_Acres", "geometry"]
    ]

    district_path = package / "04_QGIS/Groundwater/District_Potential/District_Potential.shp"
    districts = gpd.read_file(district_path)
    districts = districts.rename(
        columns={
            "district_n": "district_name",
            "coverage_p": "coverage_percent",
            "mean_gwp_s": "mean_gwp_score",
            "mean_score": "gwp_class",
        }
    )
    districts["district_key"] = districts["district_name"].map(canonical)

    contracts_path = package / "02_Data/Workbooks/Friant_Contracts_and_Entitlements.xlsx"
    entitlements = workbook_table(contracts_path, "Historical Entitlements")
    entitlements = entitlements[
        ["contractor_key", "class_1_reference_quantity_af", "class_2_reference_quantity_af", "status"]
    ].copy()
    districts = districts.merge(entitlements, how="left", left_on="district_key", right_on="contractor_key")

    district_cols = [
        "district_name", "label_id", "district_key", "coverage_percent", "mean_gwp_score", "gwp_class",
        "class_1_reference_quantity_af", "class_2_reference_quantity_af", "status", "geometry",
    ]

    class2_root = package / "Friant_Class2_Empirical_Feasibility_Package"
    class2_summary = gpd.read_file(
        class2_root / "03_GIS/Shapefiles/Class2_District_Summary/Class2_District_Summary.shp"
    )

    map_payload = {
        "study": geojson_dict(study, simplify_m=80),
        "districts": geojson_dict(districts[district_cols], simplify_m=55),
        "canals": geojson_dict(canals, simplify_m=20),
        "cropLand": geojson_dict(lulc_web[["map_group", "acres", "field_count", "geometry"]], simplify_m=50),
        "wells": geojson_dict(wells),
        "monitors": geojson_dict(monitors),
        "turnouts": geojson_dict(turnouts),
        "rise": geojson_dict(rise),
        "geology": geojson_dict(geology, simplify_m=100),
        "soils": geojson_dict(soils, simplify_m=100),
        "basins": geojson_dict(basins, simplify_m=100),
        "class2Exposure": geojson_dict(class2_summary, simplify_m=55),
    }

    raster_dir = package / "03_Results/Groundwater/Rasters"
    raster_specs = {
        "elevation": ("DEM_filled.tif", "terrain", None),
        "slope": ("Slope_degrees.tif", "magma", None),
        "drainageDensity": ("Drainage_density_km_per_km2.tif", "Blues", None),
        "streamOrder": ("Strahler_stream_order.tif", "PuBu", [0.5, 1.5, 2.5, 3.5, 4.5, 6.5]),
        "geomorphology": ("Geomorphology_code.tif", "Spectral", [0.5, 1.5, 2.5, 3.5, 4.5, 5.5, 6.5]),
        "rainfall": ("Rainfall_mean_annual_mm.tif", "YlGnBu", None),
        "groundwaterDepth": ("Depth_to_groundwater_ft.tif", "cividis_r", None),
        "recharge": ("Recharge_opportunity_index_1_5.tif", "GnBu", [0.5, 1.5, 2.5, 3.5, 4.5, 5.5]),
        "groundwaterPotential": ("Groundwater_potential_class.tif", "RdYlGn", [0.5, 1.5, 2.5, 3.5, 4.5, 5.5]),
    }
    raster_overlays = {}
    for key, (filename, cmap, bounds) in raster_specs.items():
        overlay_name = f"overlay-{key}.png"
        raster_overlays[key] = export_raster_overlay(
            raster_dir / filename, assets / overlay_name, cmap, discrete_bounds=bounds
        )
    map_payload["rasters"] = raster_overlays
    write_js(data_dir / "map-data.js", "MAP_DATA", map_payload)

    allocation = pd.read_csv(package / "03_Results/Allocation/annual_allocation_delivery_panel_1993_2025.csv")
    allocation_rows = []
    for _, row in allocation.iterrows():
        allocation_rows.append({
            "year": int(row["Year"]),
            "class1": clean_number(row["Class 1 final allocation"]),
            "class2": clean_number(row["Class 2 final allocation"]),
            "deliveryAf": clean_number(row["Usable delivery total (AF)"]),
            "deliveryRatio": clean_number(row["Delivery / benchmark"]),
            "band": row["policy_curtailment_band"],
        })

    rainfall = pd.read_csv(package / "03_Results/Groundwater/Tables/Annual_Area_Weighted_Rainfall_1993_2025.csv")
    rainfall_rows = [
        {"year": int(row.year), "mm": float(row.precipitation_mm)} for row in rainfall.itertuples()
    ]

    district_context = pd.read_csv(package / "03_Results/Allocation/district_groundwater_context_and_delivery.csv")
    fields = [
        "label_id", "district_name", "district_area_acres", "mapped_lulc_acres",
        "permanent_crop_acres", "idle_fallow_acres", "permanent_crop_share",
        "well_completion_records", "usgs_monitoring_locations", "mean_gwp_score",
        "mean_score_class", "coverage_percent", "relative_reduction_percent",
    ]
    district_rows = []
    for _, row in district_context[fields].iterrows():
        district_rows.append({key: clean_number(value) for key, value in row.items()})

    wells_annual = pd.read_csv(package / "03_Results/Allocation/district_year_well_completion_records_1993_2025.csv")
    wells_annual = wells_annual.groupby("Year", as_index=False)["well_completion_records"].sum()
    well_rows = [
        {"year": int(row.Year), "records": int(row.well_completion_records)}
        for row in wells_annual.itertuples() if pd.notna(row.Year)
    ]

    a13 = workbook_table(contracts_path, "A13 System Annual")
    a13 = a13[a13["record_type"].eq("Historical actual")]
    a13_rows = [
        {
            "year": int(row.year),
            "class1Af": clean_number(row.friant_kern_class_1_af),
            "class2Af": clean_number(row.friant_kern_class_2_af),
        }
        for row in a13.itertuples()
    ]

    crop_rows = [
        {
            "group": str(row.map_group),
            "fields": int(row.field_count),
            "acres": float(row.acres),
        }
        for row in crop_summary.itertuples()
    ]

    population_rows = [
        {"county": "Fresno", "population": 1008654},
        {"county": "Kern", "population": 909235},
        {"county": "Kings", "population": 152486},
        {"county": "Madera", "population": 156255},
        {"county": "Tulare", "population": 473117},
    ]

    episode_panel = pd.read_csv(class2_root / "02_Data/episode_panel.csv")
    primary_episodes = episode_panel[episode_panel["primary_episode"].fillna(False).astype(bool)].copy()
    class2_summary_payload = {
        "contractorYearRows": int(len(episode_panel)),
        "primaryEpisodes": int(len(primary_episodes)),
        "primaryContractors": int(primary_episodes["contractor_key"].nunique()),
        "wetYears": sorted(int(year) for year in primary_episodes["year"].dropna().unique()),
        "nextDroughtYears": sorted(
            int(year) for year in primary_episodes["next_drought_year"].dropna().unique()
        ),
        "episodesWithGroundwaterOutcomes": int(
            primary_episodes["outcome_status"].fillna("").str.contains("outcomes present").sum()
        ),
    }

    research_payload = {
        "generated": pd.Timestamp.today().date().isoformat(),
        "allocation": allocation_rows,
        "rainfall": rainfall_rows,
        "wellRecords": well_rows,
        "classDeliveries": a13_rows,
        "crops": crop_rows,
        "districts": district_rows,
        "population": population_rows,
        "class2Feasibility": class2_summary_payload,
        "counts": {
            "districts": 44,
            "lulcFields": 44066,
            "wellRecords": 10667,
            "monitoringLocations": 6698,
            "turnouts": 35,
            "canals": 2,
            "deliveryEntities": 65,
            "matchedDeliveryEntities": 33,
        },
        "notes": {
            "population": "2020 Census totals for the five counties intersecting the Friant study landscape. These are not study-area or district populations.",
            "farmland": "DWR 2023 crop/land-use field polygons. They are mapped land-use units, not legal parcels or ownership records.",
            "classReferences": "Class 1 and Class 2 spatial quantities are historical planning references and require verification against operative contracts and amendments.",
        },
    }
    write_js(data_dir / "research-data.js", "RESEARCH_DATA", research_payload)
    print(f"Prepared web data in {output}")


if __name__ == "__main__":
    main()
